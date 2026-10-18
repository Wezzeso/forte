#!/usr/bin/env python3
"""Local-only dataset identity audit and conservative group split tools.

No video pixels, audio, captions, workbook rows, or per-recording labels are
printed. Private row values stay in local data/manifests files. Source files
are opened read-only. All CLI output is aggregate JSON.
"""
from __future__ import annotations

import argparse
import concurrent.futures
import datetime as dt
import hashlib
import json
import math
import os
from pathlib import Path
import random
import re
import shutil
import subprocess
import sys
import time
import unicodedata
from collections import Counter, defaultdict
from typing import Any

VERSION = "forte-dataset-audit-v1"
SPLITS = ("train", "validation", "final_check")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def opaque_id(prefix: str, value: str, size: int = 20) -> str:
    return prefix + hashlib.sha256(value.encode("utf-8")).hexdigest()[:size]


def json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode("utf-8")


def atomic_bytes(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_bytes(payload)
    os.replace(temporary, path)


def write_json(path: Path, value: Any) -> None:
    atomic_bytes(path, json_bytes(value))


def write_jsonl(path: Path, records: list[dict]) -> None:
    atomic_bytes(path, b"".join((json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n").encode("utf-8") for record in records))


def read_jsonl(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig") as stream:
        return [json.loads(line) for line in stream if line.strip()]


def json_value(value: Any) -> Any:
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        return value.total_seconds()
    return value


def name_key(value: Any) -> str:
    """NFC is explicit; confusable letters remain different in this key."""
    return unicodedata.normalize("NFC", str(value).strip().replace("\\", "/").rsplit("/", 1)[-1])


def stem_key(value: Any) -> str:
    name = name_key(value)
    return name[:-4] if name.casefold().endswith(".mov") else name


def alias_key(value: Any) -> str:
    """Explicitly audited leading Latin X/Cyrillic Х alias, never a silent join."""
    name = stem_key(value).casefold()
    return "x" + name[1:] if name.startswith("\u0445") else name


def label_state(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    value = unicodedata.normalize("NFC", str(value or "")).strip().casefold()
    if value in {"0", "false", "no", "no_violation", "negative"} or "\u0431\u0435\u0437 \u043d\u0430\u0440\u0443\u0448" in value or "\u043d\u0435\u0442 \u043d\u0430\u0440\u0443\u0448" in value:
        return False
    if value in {"1", "true", "yes", "violation", "positive"} or "\u043d\u0430\u0440\u0443\u0448" in value:
        return True
    return None


def timestamp_sec(value: Any) -> float | None:
    if value is None or str(value).strip() in {"", "-", "\u2014"}:
        return None
    if isinstance(value, dt.time):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    if isinstance(value, dt.timedelta):
        return value.total_seconds()
    if isinstance(value, (int, float)):
        result = float(value)
        return result if math.isfinite(result) and result >= 0 else None
    text = str(value).strip().replace(",", ".")
    if re.fullmatch(r"\d+(?:\.\d+)?", text):
        return float(text)
    units = re.fullmatch(r"(?:(?P<minutes>\d+(?:\.\d+)?)\s*(?:мин[а-яё]*|min[a-z]*)\s*)?(?:(?P<seconds>\d+(?:\.\d+)?)\s*(?:сек[а-яё]*|sec[a-z]*)\s*)?", text.casefold())
    if units and any(units.groupdict().values()):
        return float(units.group("minutes") or 0) * 60 + float(units.group("seconds") or 0)
    parts = text.split(":")
    if len(parts) in (2, 3) and all(re.fullmatch(r"\d+(?:\.\d+)?", part) for part in parts):
        result = 0.0
        for part in parts:
            result = result * 60 + float(part)
        return result
    return None


def load_workbook_rows(path: Path, sheet_name: str) -> list[dict]:
    import openpyxl
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = []
        for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_number == 1 or not any(value is not None for value in values):
                continue
            if len(values) < 3:
                raise ValueError("workbook_requires_three_columns")
            rows.append({"row_id": row_number, "source_filename": str(values[0] or ""),
                         "timestamp_raw": json_value(values[1]), "label_raw": json_value(values[2]),
                         "timestamp_sec": timestamp_sec(values[1]), "has_violation": label_state(values[2])})
        return rows
    finally:
        workbook.close()


def probe_video(path: Path, ffprobe: str) -> dict:
    command = [ffprobe, "-v", "error", "-select_streams", "v:0", "-show_entries",
               "stream=index,codec_name,width,height,pix_fmt,duration,start_time,avg_frame_rate,r_frame_rate,nb_frames:stream_tags=rotate:stream_side_data=rotation:format=duration,start_time,format_name",
               "-of", "json", str(path)]
    try:
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30, check=False)
        if result.returncode:
            return {"probe_state": "failed", "probe_error": "ffprobe_nonzero"}
        payload = json.loads(result.stdout)
        streams = payload.get("streams", [])
        if not streams:
            return {"probe_state": "failed", "probe_error": "video_stream_missing"}
        video = streams[0]
        metadata = {key: video.get(key) for key in ("codec_name", "width", "height", "pix_fmt", "start_time", "avg_frame_rate", "r_frame_rate", "nb_frames")}
        duration = video.get("duration", payload.get("format", {}).get("duration"))
        metadata["duration_sec"] = float(duration) if duration not in (None, "N/A") else None
        metadata["rotation_deg"] = next((item.get("rotation") for item in video.get("side_data_list", []) if "rotation" in item), video.get("tags", {}).get("rotate", 0))
        metadata["format_name"] = payload.get("format", {}).get("format_name")
        metadata["source_format_start_time"] = payload.get("format", {}).get("start_time")
        metadata["probe_state"] = "ok"
        return metadata
    except (OSError, ValueError, subprocess.TimeoutExpired):
        return {"probe_state": "failed", "probe_error": "ffprobe_exception"}


def reconcile_rows(records: list[dict], rows: list[dict]) -> dict:
    indexes: list[tuple[str, dict[str, list[dict]]]] = []
    for strategy, transform in (("nfc_filename", name_key), ("nfc_stem", stem_key), ("explicit_x_confusable_alias", alias_key)):
        index = defaultdict(list)
        for record in records:
            index[transform(record["relative_path"])].append(record)
        indexes.append((strategy, index))
    summary = Counter()
    for row in rows:
        for strategy, index in indexes:
            transform = {"nfc_filename": name_key, "nfc_stem": stem_key, "explicit_x_confusable_alias": alias_key}[strategy]
            candidates = index.get(transform(row["source_filename"]), [])
            if not candidates:
                continue
            if len(candidates) > 1:
                row.update(join_state="ambiguous", join_strategy=strategy,
                           candidate_recording_ids=[record["recording_id"] for record in candidates])
                summary["ambiguous"] += 1
                break
            record = candidates[0]
            row.update(join_state="matched", join_strategy=strategy, recording_id=record["recording_id"],
                       content_hash=record["content_hash"], group_id=record["group_id"])
            record["workbook_row_ids"].append(row["row_id"])
            summary[strategy] += 1
            break
        else:
            row["join_state"] = "unmatched"
            summary["unmatched"] += 1
    return dict(summary)


def build_review_queue(records: list[dict], rows: list[dict]) -> list[dict]:
    row_index = {row["row_id"]: row for row in rows}
    groups = defaultdict(list)
    for record in records:
        groups[record["group_id"]].append(record)
    queue = []
    for group_id, group in sorted(groups.items()):
        linked_rows = [row_index[row_id] for record in group for row_id in record["workbook_row_ids"]]
        states = {row["has_violation"] for row in linked_rows}
        times = {row["timestamp_sec"] for row in linked_rows if row["timestamp_sec"] is not None}
        reasons = ["scenario_relationship_review_pending", "visual_annotation_review_pending"]
        if len(group) > 1:
            reasons.append("byte_identical_recordings")
        if len(states - {None}) > 1:
            reasons.append("duplicate_label_conflict")
        if len(group) > 1 and len(times) > 1:
            reasons.append("duplicate_timestamp_disagreement")
        if None in states or not linked_rows:
            reasons.append("unknown_or_missing_workbook_label")
        for record in group:
            for row_id in record["workbook_row_ids"]:
                row = row_index[row_id]
                if row["has_violation"] and row["timestamp_sec"] is None:
                    reasons.append("positive_timestamp_unparsed")
                if row["timestamp_sec"] is not None and record["duration_sec"] is not None and row["timestamp_sec"] > record["duration_sec"]:
                    reasons.append("timestamp_outside_duration")
        for record in group:
            record["review_required"] = True
            record["review_reasons"] = sorted(set(reasons))
        queue.append({"review_id": opaque_id("rev_", group_id), "group_id": group_id,
                      "content_hash": group[0]["content_hash"], "recording_ids": [record["recording_id"] for record in group],
                      "workbook_row_ids": sorted(row["row_id"] for row in linked_rows),
                      "review_state": "pending", "review_reasons": sorted(set(reasons))})
    return queue


class UnionFind:
    def __init__(self, values: list[str]):
        self.parent = {value: value for value in values}

    def find(self, value: str) -> str:
        if value not in self.parent:
            raise ValueError("scenario_group_unknown_recording")
        while self.parent[value] != value:
            self.parent[value] = self.parent[self.parent[value]]
            value = self.parent[value]
        return value

    def union(self, left: str, right: str) -> None:
        left, right = self.find(left), self.find(right)
        self.parent[max(left, right)] = min(left, right)


def grouped_records(records: list[dict], scenario: dict | None = None) -> tuple[list[list[dict]], bool]:
    union = UnionFind([record["recording_id"] for record in records])
    first = {}
    for record in records:
        content_hash = record["content_hash"]
        if content_hash in first:
            union.union(first[content_hash], record["recording_id"])
        first[content_hash] = record["recording_id"]
    reviewed_states = {"approved", "human_confirmed_proposal"}
    reviewed = bool(scenario and scenario.get("review_state") in reviewed_states)
    covered = set()
    if scenario:
        for group in scenario.get("groups", []):
            members = group.get("recording_ids", [])
            reviewed = reviewed and group.get("review_state") in reviewed_states
            for member in members:
                union.find(member)
                if member in covered:
                    raise ValueError("recording_has_multiple_scenario_memberships")
                covered.add(member)
            for member in members[1:]:
                union.union(members[0], member)
    reviewed = reviewed and covered == set(union.parent)
    groups = defaultdict(list)
    for record in records:
        groups[union.find(record["recording_id"])].append(record)
    return [sorted(value, key=lambda record: record["recording_id"]) for _, value in sorted(groups.items())], reviewed


def create_candidate(records: list[dict], rows: list[dict], seed: int, ratios: list[float], scenario: dict | None = None, excluded_final_hashes: set[str] | None = None) -> dict:
    if len(ratios) != 3 or any(value < 0 for value in ratios) or not math.isclose(sum(ratios), 1.0):
        raise ValueError("ratios_must_be_three_nonnegative_values_summing_to_one")
    by_row = {row["row_id"]: row for row in rows}
    groups, reviewed = grouped_records(records, scenario)
    scenario_membership = {recording_id: group["scenario_group_id"] for group in (scenario or {}).get("groups", []) for recording_id in group.get("recording_ids", []) if group.get("scenario_group_id")}
    strata = defaultdict(list)
    for group in groups:
        states = {by_row[row_id]["has_violation"] for record in group for row_id in record.get("workbook_row_ids", [])}
        known = states - {None}
        state = next(iter(known)) if len(known) == 1 and None not in states else None
        strata[str(state)].append((group, state))
    assignments = []
    randomizer = random.Random(seed)
    for stratum in sorted(strata):
        values = strata[stratum]
        randomizer.shuffle(values)
        # Largest-remainder allocation is deterministic; group integrity wins over counts.
        ideals = [len(values) * ratio for ratio in ratios]
        counts = [math.floor(value) for value in ideals]
        for index in sorted(range(3), key=lambda index: (-(ideals[index] - counts[index]), index))[:len(values) - sum(counts)]:
            counts[index] += 1
        eligible = [item for item in values if not any(record["content_hash"] in (excluded_final_hashes or set()) for record in item[0])]
        final_count = min(counts[2], len(eligible))
        final_values = eligible[-final_count:] if final_count else []
        final_ids = {group[0]["recording_id"] for group, _ in final_values}
        remaining = [item for item in values if item[0][0]["recording_id"] not in final_ids]
        train_count = counts[0] + counts[2] - final_count
        partitions = (remaining[:train_count], remaining[train_count:], final_values)
        for split, partition in zip(SPLITS, partitions):
            for group, state in partition:
                proposed_ids = {scenario_membership[record["recording_id"]] for record in group if record["recording_id"] in scenario_membership}
                if len(proposed_ids) > 1:
                    raise ValueError("exact_duplicates_have_conflicting_scenario_proposals")
                scenario_id = next(iter(proposed_ids)) if proposed_ids else opaque_id("scenario_", "|".join(record["recording_id"] for record in group))
                for record in group:
                    assignments.append({"sample_id": record["recording_id"], "recording_id": record["recording_id"],
                                        "content_sha256": record["content_hash"], "content_hash": record["content_hash"],
                                        "group_id": record["group_id"], "scenario_group_id": scenario_id,
                                        "split": split, "has_violation": state, "duration_sec": record.get("duration_sec")})
    assignments.sort(key=lambda record: record["sample_id"])
    assignment_digest = hashlib.sha256(json_bytes(assignments)).hexdigest()
    return {"schema_version": VERSION, "state": "proposed_unlocked", "locked": False, "approved_by": None,
            "seed": seed, "ratios": dict(zip(SPLITS, ratios)), "scenario_review_complete": reviewed,
            "training_authorized": False, "assignments_sha256": assignment_digest,
            "exposure_exclusion_applied": excluded_final_hashes is not None,
            "limitations": [] if reviewed else ["Cross-recording scenario relationships have not been approved; duplicate-only grouping is provisional."],
            "records": assignments}


def check_split(split: dict, records: list[dict], derived: list[dict] | None = None, scenario: dict | None = None, require_locked: bool = False, excluded_final_hashes: set[str] | None = None) -> dict:
    assignments = split.get("records", [])
    source = {record["recording_id"]: record for record in records}
    by_id = {}
    errors = Counter()
    errors["duplicate_source_recording_id"] = len(records) - len(source)
    content_splits = defaultdict(set)
    group_splits = defaultdict(set)
    for assignment in assignments:
        recording_id = assignment.get("recording_id", assignment.get("sample_id"))
        if recording_id in by_id:
            errors["duplicate_recording_assignment"] += 1
        by_id[recording_id] = assignment
        if recording_id not in source:
            errors["unknown_recording_assignment"] += 1
            continue
        record = source[recording_id]
        if assignment.get("sample_id", recording_id) != recording_id:
            errors["assignment_identity_alias_mismatch"] += 1
        if assignment.get("split") not in SPLITS:
            errors["invalid_split"] += 1
        if assignment.get("content_sha256", assignment.get("content_hash")) != record["content_hash"] or assignment.get("content_hash", assignment.get("content_sha256")) != record["content_hash"]:
            errors["content_hash_mismatch"] += 1
        if assignment.get("group_id") != record["group_id"]:
            errors["assignment_group_mismatch"] += 1
        content_splits[record["content_hash"]].add(assignment.get("split"))
        group_splits[record["group_id"]].add(assignment.get("split"))
    errors["unassigned_recording"] = len(set(source) - set(by_id))
    errors["content_crosses_split"] = sum(len(values) > 1 for values in content_splits.values())
    errors["recording_group_crosses_split"] = sum(len(values) > 1 for values in group_splits.values())
    groups, reviewed = grouped_records(records, scenario)
    errors["scenario_group_crosses_split"] = sum(len({by_id[record["recording_id"]]["split"] for record in group if record["recording_id"] in by_id}) > 1 for group in groups)
    errors["exposed_scenario_in_final_check"] = sum(any(record["content_hash"] in (excluded_final_hashes or set()) for record in group) and any(by_id.get(record["recording_id"], {}).get("split") == "final_check" for record in group) for group in groups)
    if "assignments_sha256" in split and split["assignments_sha256"] != hashlib.sha256(json_bytes(assignments)).hexdigest():
        errors["assignment_digest_mismatch"] += 1
    if require_locked and not (split.get("locked") is True and split.get("state") == "locked" and split.get("approved_by") and split.get("training_authorized") is True and reviewed):
        errors["split_not_approved_for_training"] += 1
    if require_locked and excluded_final_hashes is None:
        errors["exposure_inventory_missing"] += 1
    for item in derived or []:
        parent_id = item.get("source_recording_id", item.get("recording_id"))
        assignment = by_id.get(parent_id)
        if assignment is None:
            errors["derived_unknown_parent"] += 1
            continue
        if item.get("split") != assignment["split"]:
            errors["derived_split_mismatch"] += 1
        if item.get("group_id") != assignment["group_id"]:
            errors["derived_group_mismatch"] += 1
        if item.get("content_hash", item.get("content_sha256")) != assignment["content_hash"]:
            errors["derived_content_mismatch"] += 1
    nonzero = {key: value for key, value in sorted(errors.items()) if value}
    return {"ok": not nonzero, "recordings": len(source), "assignments": len(assignments), "derived_records": len(derived or []),
            "errors": nonzero, "scenario_review_complete": reviewed, "locked": bool(split.get("locked")),
            "split_counts": dict(Counter(record.get("split") for record in assignments))}


def audit(args: argparse.Namespace) -> dict:
    started = time.perf_counter()
    root = Path(args.root).resolve()
    video_root = (root / args.videos).resolve()
    workbook = (root / args.workbook).resolve()
    output = (root / args.output).resolve()
    split_root = (root / args.splits).resolve()
    ffprobe = args.ffprobe or shutil.which("ffprobe")
    if not ffprobe:
        raise ValueError("ffprobe_required")
    files = sorted((path for path in video_root.rglob("*") if path.is_file() and path.suffix.casefold() == ".mov"), key=lambda path: path.relative_to(root).as_posix())
    if not files:
        raise ValueError("no_video_files")
    workbook_hash = sha256_file(workbook)
    initial = {path: (path.stat().st_size, path.stat().st_mtime_ns) for path in files}
    def inspect(path: Path) -> dict:
        relative = path.relative_to(root).as_posix()
        content_hash = sha256_file(path)
        metadata = probe_video(path, str(ffprobe))
        group_id = "grp_" + content_hash[:20]
        return {"recording_id": opaque_id("rec_", relative), "sample_id": opaque_id("rec_", relative),
                "content_hash": content_hash, "content_sha256": content_hash, "content_id": "vid_" + content_hash[:20],
                "recording_group_id": group_id, "group_id": group_id, "scenario_group_id": None,
                "relative_path": relative, "source_path": relative, "split": "unassigned",
                "duration_sec": metadata.get("duration_sec"), "width": metadata.get("width"), "height": metadata.get("height"),
                "workbook_row_ids": [], "metadata": {**metadata, "size_bytes": initial[path][0], "mtime_ns": initial[path][1]}}
    with concurrent.futures.ThreadPoolExecutor(max_workers=args.workers) as pool:
        records = list(pool.map(inspect, files))
    rows = load_workbook_rows(workbook, args.sheet)
    joins = reconcile_rows(records, rows)
    queue = build_review_queue(records, rows)
    unchanged = all(initial[path] == (path.stat().st_size, path.stat().st_mtime_ns) and sha256_file(path) == record["content_hash"] for path, record in zip(files, records)) and sha256_file(workbook) == workbook_hash
    if not unchanged:
        raise ValueError("source_changed_during_audit")
    core_validation = "unavailable"
    try:
        sys.path.insert(0, str(root / "src"))
        from forte.shared.schemas import RecordingRecord
        for record in records:
            RecordingRecord.model_validate({key: record[key] for key in ("recording_id", "content_hash", "recording_group_id", "relative_path", "scenario_group_id", "split", "duration_sec", "width", "height", "workbook_row_ids", "metadata")})
        core_validation = "passed"
    except ImportError:
        pass
    scenario_path = split_root / "scenario_groups.review.json"
    if not scenario_path.exists():
        write_json(scenario_path, {"schema_version": VERSION, "review_state": "pending", "reviewer": None,
                                  "notes": "Exact-content grouping only. A local reviewer must review scenario relationships across all recordings before approval.",
                                  "groups": [{"scenario_group_id": opaque_id("scenario_", item["group_id"]), "recording_ids": item["recording_ids"], "review_state": "pending"} for item in queue]})
    scenario = json.loads(scenario_path.read_text(encoding="utf-8-sig"))
    write_jsonl(output / "recordings.jsonl", records)
    write_jsonl(output / "workbook_rows.jsonl", rows)
    write_jsonl(output / "review_queue.jsonl", queue)
    registry_path = root / "benchmarks/private/model-exposure.jsonl"
    exposure_hashes = None
    if registry_path.exists():
        from exposure_registry import read_registry
        exposure_hashes = read_registry(registry_path, records)["content_hashes"]
    candidate = create_candidate(records, rows, args.seed, args.ratios, scenario, exposure_hashes)
    candidate["source_manifest_sha256"] = sha256_file(output / "recordings.jsonl")
    candidate["source_workbook_sha256"] = workbook_hash
    candidate["scenario_groups_sha256"] = sha256_file(scenario_path)
    write_json(split_root / "candidate_split.json", candidate)
    split_check = check_split(candidate, records, scenario=scenario, excluded_final_hashes=exposure_hashes)
    write_json(split_root / "integrity_report.json", split_check)
    counts = Counter(record["content_hash"] for record in records)
    durations = [record["duration_sec"] for record in records if record["duration_sec"] is not None]
    audit_ok = len(files) == len(rows) and not joins.get("ambiguous") and not joins.get("unmatched") and all(record["workbook_row_ids"] and record["metadata"]["probe_state"] == "ok" for record in records)
    summary = {"ok": audit_ok, "schema_version": VERSION, "generated_at_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
               "recordings": len(records), "unique_contents": len(counts), "duplicate_groups": sum(count > 1 for count in counts.values()),
               "redundant_recordings": sum(count - 1 for count in counts.values()), "workbook_rows": len(rows),
               "workbook_sha256": workbook_hash, "join_counts": joins,
               "unlinked_recordings": sum(not record["workbook_row_ids"] for record in records),
               "multiple_rows_per_recording": sum(len(record["workbook_row_ids"]) > 1 for record in records),
               "label_counts": dict(Counter("positive" if row["has_violation"] is True else "negative" if row["has_violation"] is False else "unknown" for row in rows)),
               "metadata_probe_counts": dict(Counter(record["metadata"]["probe_state"] for record in records)),
               "codec_counts": dict(Counter(record["metadata"].get("codec_name", "unknown") for record in records)),
               "resolution_counts": dict(Counter(f"{record['width']}x{record['height']}" for record in records)),
               "duration_sec": {"min": min(durations, default=None), "max": max(durations, default=None), "sum": round(sum(durations), 3)},
               "review_reason_group_counts": dict(Counter(reason for item in queue for reason in item["review_reasons"])),
               "source_bytes_unchanged": unchanged, "shared_schema_validation": core_validation,
               "split_integrity": split_check, "elapsed_sec": round(time.perf_counter() - started, 3),
               "privacy": "Metadata/hash audit only; media and private row values were not printed."}
    write_json(output / "audit_summary.json", summary)
    return summary


def self_test() -> dict:
    checks = 0
    assert name_key("X-synthetic-1.mov") != name_key("\u0425-synthetic-1.mov"); checks += 1
    assert alias_key("X-synthetic-1.mov") == alias_key("\u0425-synthetic-1.mov"); checks += 1
    assert timestamp_sec("01:02.5") == 62.5; checks += 1
    assert timestamp_sec("00:01:02") == 62; checks += 1
    assert timestamp_sec("1минута 2 секунда") == 62; checks += 1
    assert timestamp_sec("3 секунда") == 3; checks += 1
    assert label_state("\u0431\u0435\u0437 \u043d\u0430\u0440\u0443\u0448\u0435\u043d\u0438\u044f") is False; checks += 1
    records = [{"recording_id": f"rec_{index}", "content_hash": f"hash_{index // 2}", "group_id": f"grp_{index // 2}", "relative_path": f"videos/X-synthetic-{index}.mov", "duration_sec": 10, "workbook_row_ids": [index + 2]} for index in range(12)]
    rows = [{"row_id": index + 2, "has_violation": bool((index // 2) % 2)} for index in range(12)]
    split = create_candidate(records, rows, 17, [0.7, 0.15, 0.15])
    assert split == create_candidate(records, rows, 17, [0.7, 0.15, 0.15]); checks += 1
    assert check_split(split, records)["ok"]; checks += 1
    assert not check_split(split, records, require_locked=True)["ok"]; checks += 1
    corrupted = json.loads(json.dumps(split))
    corrupted["records"][0]["split"] = "validation" if corrupted["records"][1]["split"] != "validation" else "train"
    assert not check_split(corrupted, records)["ok"]; checks += 1
    corrupted = json.loads(json.dumps(split))
    corrupted["records"][0]["content_hash"] = "changed_hash"
    assert not check_split(corrupted, records)["ok"]; checks += 1
    assignment = split["records"][0]
    child = {"source_recording_id": assignment["recording_id"], "split": assignment["split"], "group_id": assignment["group_id"], "content_hash": assignment["content_hash"]}
    assert check_split(split, records, [child])["ok"]; checks += 1
    child["split"] = "invalid"
    assert not check_split(split, records, [child])["ok"]; checks += 1
    ambiguous_records = [{"recording_id": "a", "relative_path": "a/X-synthetic-1.mov", "workbook_row_ids": []}, {"recording_id": "b", "relative_path": "b/X-synthetic-1.mov", "workbook_row_ids": []}]
    ambiguous_rows = [{"row_id": 2, "source_filename": "X-synthetic-1.mov"}]
    assert reconcile_rows(ambiguous_records, ambiguous_rows) == {"ambiguous": 1}; checks += 1
    return {"ok": True, "synthetic_checks_passed": checks}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    audit_parser = sub.add_parser("audit")
    audit_parser.add_argument("--root", default=str(Path(__file__).resolve().parents[1]))
    audit_parser.add_argument("--videos", default="videos")
    audit_parser.add_argument("--workbook", default="open_labels.xlsx")
    audit_parser.add_argument("--sheet", default="labels")
    audit_parser.add_argument("--output", default="data/manifests")
    audit_parser.add_argument("--splits", default="data/splits")
    audit_parser.add_argument("--ffprobe")
    audit_parser.add_argument("--workers", type=int, default=4)
    audit_parser.add_argument("--seed", type=int, default=20261017)
    audit_parser.add_argument("--ratios", nargs=3, type=float, default=[0.7, 0.15, 0.15])
    check_parser = sub.add_parser("check-split")
    check_parser.add_argument("--manifest", type=Path, default=Path("data/manifests/recordings.jsonl"))
    check_parser.add_argument("--split", type=Path, default=Path("data/splits/candidate_split.json"))
    check_parser.add_argument("--scenario-groups", type=Path, default=Path("data/splits/scenario_groups.review.json"))
    check_parser.add_argument("--workbook", type=Path, default=Path("open_labels.xlsx"))
    check_parser.add_argument("--workbook-rows", type=Path, default=Path("data/manifests/workbook_rows.jsonl"))
    check_parser.add_argument("--exposure-registry", type=Path, default=Path("benchmarks/private/model-exposure.jsonl"))
    check_parser.add_argument("--derived", type=Path)
    check_parser.add_argument("--require-locked", action="store_true")
    sub.add_parser("self-test")
    args = parser.parse_args()
    try:
        if args.command == "audit":
            result = audit(args)
        elif args.command == "check-split":
            scenario = json.loads(args.scenario_groups.read_text(encoding="utf-8-sig")) if args.scenario_groups.exists() else None
            split = json.loads(args.split.read_text(encoding="utf-8-sig"))
            records = read_jsonl(args.manifest)
            registry = None
            if args.exposure_registry.exists():
                from exposure_registry import read_registry
                registry = read_registry(args.exposure_registry, records)
            result = check_split(split, records, read_jsonl(args.derived) if args.derived else None, scenario, args.require_locked, registry["content_hashes"] if registry else None)
            if split.get("exposure_registry_tip_sha256") and (registry is None or split["exposure_registry_tip_sha256"] not in {event["event_sha256"] for event in registry["events"]}):
                result["errors"]["exposure_registry_snapshot_missing"] = 1
                result["ok"] = False
            for field, path in (("source_manifest_sha256", args.manifest), ("source_workbook_sha256", args.workbook), ("source_workbook_rows_sha256", args.workbook_rows), ("scenario_groups_sha256", args.scenario_groups)):
                if field in split and (not path.is_file() or sha256_file(path) != split[field]):
                    result["errors"][field + "_mismatch"] = 1
                    result["ok"] = False
        else:
            result = self_test()
        print(json.dumps(result, ensure_ascii=True, sort_keys=True))
        return 0 if result.get("ok", True) else 2
    except Exception as exc:
        # Do not leak paths/row values from third-party exception messages.
        code = str(exc) if isinstance(exc, ValueError) and re.fullmatch("[a-z_]+", str(exc)) else "operation_failed"
        print(json.dumps({"ok": False, "error_type": type(exc).__name__, "error_code": code}))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
